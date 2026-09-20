#!/usr/bin/env python3
"""
combine_certs.py

Combines a folder/zip of many single-file Excel exports (Quantity, Line, Product,
Description, Confirmation Notes) into one workbook, filtered to specific Product
codes, with a blank separator row before each source file's block. Optionally
merges in a "LIST" CSV whose BLANK column supplies a certificate-reference note
that gets slotted into the Confirmation Notes cell of each separator row, in order.

Usage:
    python combine_certs.py FILES.zip [--list LIST.csv] [--out OUTPUT.xlsx]

    FILES.zip   Zip archive containing the individual .xlsx files (any nesting,
                any __MACOSX junk is ignored automatically).
    --list      Optional CSV with columns including Number, Customer\\nRef (or
                "Customer Ref"), Job Number (Particulars), and BLANK. If given,
                the BLANK value from list row N is written into the Confirmation
                Notes cell of separator row N (matched by ORDER only -- the two
                inputs must already be in the same order and have the same count
                of files / list rows).
    --out       Output path (default: Combined_Filtered_Products.xlsx next to
                the input zip).

Business rules encoded here (confirmed with the user across several rounds --
change these constants/functions if the rules ever change):

  1. PRODUCT FILTER -- a row is kept only if its Product value:
       - starts with "ISO4017G", OR
       - starts with "933CEASS", OR
       - contains "HSFG" anywhere.

  2. FILE NUMBERING -- the file's identifying number is the first run of digits
     found in its filename (e.g. "01.xlsx" -> 1, "Excel (79).xlsx" -> 79,
     "Excel-23.xlsx" -> 23, "Excel.xlsx" (no digits) -> 1). Files are combined
     in ascending numeric order and labelled as a zero-padded 2-digit string
     ("01".."99") in the File column.

  3. LAYOUT -- one blank row is inserted immediately BEFORE every file's block
     of rows, including before the very first file.

  4. BLANK / CERT-NOTE FORMULA (from the LIST csv), for reference / regeneration:
       BLANK = UPPER(
                 "NEW CERTS " + Number
                 + (" PO " + CustomerRef   if CustomerRef is non-empty, else "")
                 + (" REF " + JobNumber    if JobNumber   is non-empty, else "")
               )
     Before inserting, strip a leading "PO" from CustomerRef (and a leading
     "REF" from JobNumber), including any immediate separator (space/dash/
     colon), so the result never contains "PO PO ..." or "REF REF ...".
     CustomerRef/JobNumber are NOT trimmed of internal/trailing whitespace
     otherwise -- only the redundant leading prefix is stripped.

  5. MERGE WITH LIST -- the BLANK values are matched to separator rows purely
     by POSITION/ORDER (list row 1 -> 1st separator row, list row 2 -> 2nd,
     etc.), not by any key. The user has confirmed the order will always
     correspond. The count of list rows must equal the count of files /
     separator rows, or this script raises an error rather than guessing.
     The value is written into the Confirmation Notes column (not File).
"""

import argparse
import csv
import re
import sys
import zipfile
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill

HEADERS = ["File", "Quantity", "Line", "Product", "Description", "Confirmation Notes"]

HEADER_FONT = Font(name="Arial", bold=True, color="FFFFFF")
HEADER_FILL = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
BODY_FONT = Font(name="Arial")
FILE_FONT = Font(name="Arial", bold=True)
NOTE_FONT = Font(name="Arial", bold=True, italic=True, color="C00000")


# ---------------------------------------------------------------------------
# Rule 1: product filter
# ---------------------------------------------------------------------------
def product_matches(product) -> bool:
    if product is None:
        return False
    p = str(product).strip()
    return p.startswith("ISO4017G") or p.startswith("933CEASS") or "HSFG" in p


# ---------------------------------------------------------------------------
# Rule 2: file numbering
# ---------------------------------------------------------------------------
def file_number(filename: str) -> int:
    base = Path(filename).stem
    m = re.search(r"(\d+)", base)
    return int(m.group(1)) if m else 1


# ---------------------------------------------------------------------------
# Rule 4: BLANK / cert-note text builder
# ---------------------------------------------------------------------------
def _strip_leading(value: str, prefix: str) -> str:
    """Strip a leading prefix (case-insensitive) plus any immediate separator."""
    if value is None:
        return value
    stripped = value.lstrip()
    if stripped.upper().startswith(prefix.upper()):
        rest = stripped[len(prefix):].lstrip(" -:")
        return rest
    return value


def build_blank(number: str, customer_ref: str, job_number: str) -> str:
    out = f"NEW CERTS {number}"
    if customer_ref:
        out += f" PO {_strip_leading(customer_ref, 'PO')}"
    if job_number:
        out += f" REF {_strip_leading(job_number, 'REF')}"
    return out.upper()


# ---------------------------------------------------------------------------
# Zip extraction (ignores __MACOSX and non-.xlsx junk)
# ---------------------------------------------------------------------------
def extract_xlsx_from_zip(zip_path: Path, dest_dir: Path):
    with zipfile.ZipFile(zip_path) as zf:
        names = [
            n for n in zf.namelist()
            if n.lower().endswith(".xlsx") and "__MACOSX" not in n and not Path(n).name.startswith("._")
        ]
        for n in names:
            zf.extract(n, dest_dir)
    return sorted(dest_dir.rglob("*.xlsx"), key=lambda p: file_number(p.name))


# ---------------------------------------------------------------------------
# Reading the LIST csv
# ---------------------------------------------------------------------------
def read_list_csv(list_path: Path):
    """Returns a list of BLANK strings, one per data row, in file order.

    Uses the BLANK column if present; otherwise recomputes it from
    Number / Customer Ref / Job Number (Particulars) using build_blank().
    """
    with open(list_path, newline="", encoding="utf-8-sig") as f:
        reader = csv.DictReader(f)
        fieldnames = reader.fieldnames or []
        # Column names can vary slightly (e.g. "Customer\nRef" vs "Customer Ref")
        def find_col(*candidates):
            for c in candidates:
                for fn in fieldnames:
                    if fn.replace("\n", " ").strip().lower() == c.lower():
                        return fn
            return None

        col_number = find_col("Number")
        col_custref = find_col("Customer Ref")
        col_jobnum = find_col("Job Number (Particulars)")
        col_blank = find_col("BLANK")

        results = []
        for row in reader:
            if col_blank and row.get(col_blank, "").strip() != "":
                results.append(row[col_blank])
            else:
                results.append(
                    build_blank(
                        row.get(col_number, ""),
                        row.get(col_custref, ""),
                        row.get(col_jobnum, ""),
                    )
                )
        return results


# ---------------------------------------------------------------------------
# Main combine logic
# ---------------------------------------------------------------------------
def combine(zip_path: Path, list_path: Path | None, out_path: Path):
    work_dir = out_path.parent / (out_path.stem + "_extracted_tmp")
    work_dir.mkdir(parents=True, exist_ok=True)

    files = extract_xlsx_from_zip(zip_path, work_dir)
    if not files:
        raise SystemExit(f"No .xlsx files found inside {zip_path}")

    wb_out = openpyxl.Workbook()
    ws_out = wb_out.active
    ws_out.title = "Combined"
    ws_out.append(HEADERS)
    for cell in ws_out[1]:
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center")

    separator_rows = []  # track row numbers of the blank rows, in order

    for f in files:
        n = file_number(f.name)
        label = f"{n:02d}"

        wb_in = openpyxl.load_workbook(f, data_only=True)
        ws_in = wb_in.worksheets[0]
        rows = list(ws_in.iter_rows(values_only=True))

        matched = []
        for row in rows[1:]:
            if not row or not any(c is not None and str(c).strip() != "" for c in row):
                continue
            quantity, line, product, description, notes = (list(row) + [None] * 5)[:5]
            if product_matches(product):
                matched.append((quantity, line, product, description, notes))

        # Rule 3: blank row before every file, including the first
        ws_out.append([None] * len(HEADERS))
        separator_rows.append(ws_out.max_row)

        for r in matched:
            ws_out.append([label, *r])
            for cell in ws_out[ws_out.max_row]:
                cell.font = BODY_FONT
            ws_out.cell(row=ws_out.max_row, column=1).font = FILE_FONT

    # Rule 5: merge in LIST BLANK values by position, into Confirmation Notes
    if list_path is not None:
        blanks = read_list_csv(list_path)
        if len(blanks) != len(separator_rows):
            raise SystemExit(
                f"Row-count mismatch: {len(blanks)} list rows vs "
                f"{len(separator_rows)} files/separator rows. "
                "Refusing to guess a mapping -- fix the inputs so counts match."
            )
        conf_col = HEADERS.index("Confirmation Notes") + 1
        for row_num, text in zip(separator_rows, blanks):
            c = ws_out.cell(row=row_num, column=conf_col, value=text)
            c.font = NOTE_FONT

    widths = [10, 12, 8, 22, 40, 45]
    for i, w in enumerate(widths, start=1):
        ws_out.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
    ws_out.freeze_panes = "A2"

    wb_out.save(out_path)

    # cleanup extracted tmp dir
    import shutil
    shutil.rmtree(work_dir, ignore_errors=True)

    return out_path, len(files), len(separator_rows)


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("zip_path", type=Path, help="Zip file containing the individual .xlsx files")
    ap.add_argument("--list", dest="list_path", type=Path, default=None, help="Optional LIST csv to merge in")
    ap.add_argument("--out", dest="out_path", type=Path, default=None, help="Output .xlsx path")
    args = ap.parse_args()

    out_path = args.out_path or (args.zip_path.parent / "Combined_Filtered_Products.xlsx")
    out_path, n_files, n_sep = combine(args.zip_path, args.list_path, out_path)
    print(f"Saved: {out_path}")
    print(f"Files combined: {n_files}")
    print(f"Separator rows: {n_sep}")


if __name__ == "__main__":
    main()
